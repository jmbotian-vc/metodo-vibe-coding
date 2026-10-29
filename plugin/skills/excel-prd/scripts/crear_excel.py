"""Crea el Excel de PRD y seguimiento de un proyecto con la estructura de GF-CRM.

Uso:  python3 crear_excel.py <ruta_salida.xlsx> "<Nombre del proyecto>" [--ruta A|B]
Requiere openpyxl (pip install openpyxl).
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = "2E5FA3"
GRIS = "F2F2F2"

ESTADOS_TAREA = "Por hacer,En curso,En revisión,Hecha,Bloqueada"
ESTADOS_REQ = "Propuesta,Aprobado,En curso,Hecho,Descartado"
SI_NO = "Sí,No"

# pestaña: (descripción, columnas, {columna: lista desplegable})
PESTANAS = {
    "PRD": ("Cada fila es un requisito verificable. No se marca como hecho sin evidencia en CHECKLIST o CHANGELOG.",
            ["REQ-ID", "Épica", "Módulo", "Requisito", "Usuario / rol", "Prioridad", "Dependencias",
             "Criterio de aceptación", "Estado", "Check", "Evidencia / nota"],
            {"Prioridad": "Alta,Media,Baja", "Estado": ESTADOS_REQ, "Check": SI_NO}),
    "BACKLOG": ("Cada tarea termina con su check de aceptación y referencia a un requisito. Hecha = PR mergeado y enlazado.",
                ["TASK-ID", "REQ-ID", "Fase", "Módulo", "Tarea", "Responsable", "Estado", "Prioridad",
                 "Criterio de terminado", "Check", "Evidencia / commit", "Fecha objetivo"],
                {"Estado": ESTADOS_TAREA, "Prioridad": "Alta,Media,Baja", "Check": SI_NO}),
    "ROADMAP": ("Fases del desarrollo (los milestones). Las fechas son objetivos y se ajustan desde BACKLOG.",
                ["Fase", "Objetivo", "Inicio objetivo", "Fin objetivo", "Entregable", "Dependencias",
                 "Estado", "Check", "Criterio de salida"],
                {"Estado": ESTADOS_TAREA, "Check": SI_NO}),
    "CHECKLIST": ("Una tarea no se cierra hasta que todos sus checks tengan evidencia.",
                  ["CHECK-ID", "TASK-ID", "Área", "Comprobación", "Resultado esperado", "Ejecutado por",
                   "Fecha", "Estado", "Evidencia / enlace", "Observación"],
                  {"Estado": "Pendiente,OK,Falla"}),
    "CHANGELOG": ("Registro obligatorio de cambios: no se considera aplicado si no tiene fila con check, evidencia y versión.",
                  ["CHANGE-ID", "Fecha", "Tipo", "Área", "Descripción del cambio", "Motivo", "Solicitado por",
                   "Ejecutado por", "Versión / commit", "Estado", "Check", "Evidencia / prueba", "Impacto"],
                  {"Tipo": "Nueva función,Mejora,Corrección,Seguridad,Infraestructura,Documentación",
                   "Check": SI_NO}),
    "DECISIONES": ("Una decisión evita reinterpretaciones posteriores y protege el alcance.",
                   ["DEC-ID", "Fecha", "Tema", "Decisión", "Alternativas descartadas", "Razón", "Aprobador",
                    "Estado", "Check"],
                   {"Check": SI_NO}),
    "RIESGOS": ("Un riesgo abierto debe tener responsable y próxima acción.",
                ["RISK-ID", "Riesgo", "Probabilidad", "Impacto", "Mitigación", "Estado", "Responsable",
                 "Próxima acción", "Check"],
                {"Probabilidad": "Alta,Media,Baja", "Impacto": "Alto,Medio,Bajo", "Estado": "Abierto,Mitigado,Cerrado"}),
    "ROLES": ("Permisos por rol. Se verifican en el servidor; ocultar botones no sustituye la autorización.",
              ["Rol", "Usuario tipo", "Puede consultar", "Puede crear / editar", "Puede revisar",
               "Puede administrar usuarios", "Regla"], {}),
    "DATOS": ("Mapa de datos: tablas o pestañas, uso, lectura/escritura y privacidad.",
              ["Fuente", "Tabla / pestaña", "Uso", "Lectura", "Escritura", "Datos personales", "Privacidad / nota"],
              {"Datos personales": SI_NO}),
    "API": ("Cada función pública valida identidad, rol y datos de entrada.",
            ["API-ID", "Función", "Módulo", "Lectura", "Escritura", "Autorización", "Respuesta", "Estado",
             "Prueba mínima", "Check"], {"Check": SI_NO}),
    "INCIDENTES": ("Plan de incidentes: cómo cortar, rotar, restaurar y avisar por cada servicio.",
                   ["Servicio", "Cómo cortar el acceso", "Claves a rotar y dónde", "Cómo restaurar",
                    "A quién avisar", "Última revisión"], {}),
}


def cabecera(ws, titulo, descripcion, columnas):
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"] = descripcion
    ws["A2"].font = Font(italic=True, color="666666")
    for i, col in enumerate(columnas, start=1):
        c = ws.cell(row=4, column=i, value=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = max(14, min(40, len(col) + 6))
    ws.freeze_panes = "A5"


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    salida, nombre = sys.argv[1], sys.argv[2]
    ruta = sys.argv[sys.argv.index("--ruta") + 1] if "--ruta" in sys.argv else ""

    wb = Workbook()
    ini = wb.active
    ini.title = "INICIO"
    ini["A1"] = f"{nombre} · PRD, backlog y planificación"
    ini["A1"].font = Font(bold=True, size=16, color=AZUL)
    ini["A2"] = f"Documento vivo de requisitos, desarrollo y control de cambios · Versión 0.1 · {date.today().isoformat()}"
    ini["A2"].font = Font(italic=True, color="666666")
    filas = [
        ("Propósito", ""), ("Para quién", ""), ("Ruta", {"A": "A · Apps Script + Google Sheets + GitHub",
                                                       "B": "B · Next.js + Convex + GitHub + Railway"}.get(ruta, "")),
        ("Repositorio", ""), ("URL de pruebas", ""), ("URL de producción", ""),
    ]
    for i, (k, v) in enumerate(filas, start=4):
        ini.cell(row=i, column=1, value=k).font = Font(bold=True)
        ini.cell(row=i, column=2, value=v)
    ini.cell(row=11, column=1, value="Indicador").font = Font(bold=True)
    ini.cell(row=11, column=2, value="Valor").font = Font(bold=True)
    indicadores = [
        ("Requisitos", '=COUNTA(PRD!A5:A1000)'),
        ("Requisitos hechos", '=COUNTIF(PRD!I5:I1000,"Hecho")'),
        ("Tareas", '=COUNTA(BACKLOG!A5:A1000)'),
        ("Tareas hechas", '=COUNTIF(BACKLOG!G5:G1000,"Hecha")'),
        ("Tareas en curso", '=COUNTIF(BACKLOG!G5:G1000,"En curso")'),
        ("Riesgos abiertos", '=COUNTIF(RIESGOS!F5:F1000,"Abierto")'),
    ]
    for i, (k, f) in enumerate(indicadores, start=12):
        ini.cell(row=i, column=1, value=k)
        ini.cell(row=i, column=2, value=f)
    ini.column_dimensions["A"].width = 22
    ini.column_dimensions["B"].width = 60

    for titulo, (desc, cols, listas) in PESTANAS.items():
        ws = wb.create_sheet(titulo)
        cabecera(ws, f"{titulo} · {desc.split('.')[0]}", desc, cols)
        for col, opciones in listas.items():
            letra = ws.cell(row=4, column=cols.index(col) + 1).column_letter
            dv = DataValidation(type="list", formula1=f'"{opciones}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letra}5:{letra}1000")
        ws.sheet_properties.tabColor = AZUL if titulo in ("PRD", "BACKLOG") else "999999"

    wb.save(salida)
    print(f"Creado {salida}")


if __name__ == "__main__":
    main()
